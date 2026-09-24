# -*- coding: utf-8 -*-
"""
Capa Silver: Módulo de Novedades / Referidos para TI (Altas y Bajas).
Basado en el proceso de Lina Builes para el área de Tecnología.
"""
import sys
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

import logging
import pandas as pd
from datetime import datetime
from config.settings import SILVER_DIR, OUTPUTS_DIR

logger = logging.getLogger("SilverTINovedades")
logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")


class SilverTINovedadesProcessor:
    def __init__(self):
        self.silver_dir = SILVER_DIR
        self.outputs_dir = OUTPUTS_DIR

    def generate_ti_template(self, year: int, month: int) -> pd.DataFrame:
        logger.info(f"Generando plantilla de novedades para TI ({year}-{month:02d})...")
        self.outputs_dir.mkdir(parents=True, exist_ok=True)
        self.silver_dir.mkdir(parents=True, exist_ok=True)

        df_emp = None
        try:
            from src.database.azure_data_service import azure_data_service
            if azure_data_service.is_connected():
                df_emp = azure_data_service.get_silver_employees()
        except Exception as e:
            logger.warning(f"Error cargando silver employees desde Azure: {e}")
            df_emp = None

        if df_emp is None or df_emp.empty:
            emp_parquet = self.silver_dir / "silver_employees.parquet"
            if emp_parquet.exists():
                df_emp = pd.read_parquet(emp_parquet)
            else:
                try:
                    from src.processing.silver_employees import SilverEmployeesProcessor
                    df_emp = SilverEmployeesProcessor().process()
                except Exception as e:
                    logger.error(f"Error procesando silver employees: {e}")
                    raise FileNotFoundError("No se encontró silver_employees ni en Azure ni local.")

        df_emp = df_emp.copy()
        df_emp["fecha_ingreso_str"] = df_emp["fecha_ingreso"].fillna("").astype(str).str[:10]
        df_emp["fecha_retiro_str"] = df_emp["fecha_retiro"].fillna("").astype(str).str[:10]

        start_month = f"{year}-{month:02d}-01"
        import calendar
        _, last_day = calendar.monthrange(year, month)
        end_month = f"{year}-{month:02d}-{last_day:02d}"

        # Ingresos del mes
        df_ingresos = df_emp[
            (df_emp["fecha_ingreso_str"] >= start_month) & 
            (df_emp["fecha_ingreso_str"] <= end_month)
        ].copy()
        df_ingresos["movimiento"] = "INGRESO"
        df_ingresos["fecha_movimiento"] = df_ingresos["fecha_ingreso_str"]

        # Retiros del mes
        df_retiros = df_emp[
            (df_emp["fecha_retiro_str"] != "") &
            (df_emp["fecha_retiro_str"] != "None") &
            (df_emp["fecha_retiro_str"] != "nan") &
            (df_emp["fecha_retiro_str"] >= start_month) & 
            (df_emp["fecha_retiro_str"] <= end_month)
        ].copy()
        df_retiros["movimiento"] = "RETIRO"
        df_retiros["fecha_movimiento"] = df_retiros["fecha_retiro_str"]

        # Consolidar
        df_novedades = pd.concat([df_ingresos, df_retiros], ignore_index=True)
        df_novedades = df_novedades.drop_duplicates(subset=["documento", "movimiento"])

        # Formato exacto que pide el área de TI
        ti_records = []
        for _, r in df_novedades.iterrows():
            ti_records.append({
                "Tipo de Doc": r["tipo_documento"],
                "Número": r["documento"],
                "(*) Nombre": r["nombres"],
                "(*) Apellido": f"{r['primer_apellido']} {r['segundo_apellido']}".strip(),
                "picture_url": r.get("picture_url", ""),
                "CORREO": r["email_corporativo"] if r["email_corporativo"] else f"{r['nombres'].lower().replace(' ', '')}@maaji.co",
                "MOVIMIENTO": r["movimiento"],
                "FECHA NOVEDAD": r["fecha_movimiento"],
                "CARGO": r["cargo_nombre"],
                "EMPRESA": r["empresa_nombre"],
                "SEDE": r["sede"]
            })

        df_ti = pd.DataFrame(ti_records)
        
        # Merge manual novelties if any exist
        manual_path = self.silver_dir / "ti_manual_novedades.json"
        if manual_path.exists():
            try:
                import json
                with open(manual_path, "r", encoding="utf-8") as f:
                    manual_items = json.load(f)
                if manual_items:
                    df_manual = pd.DataFrame(manual_items)
                    df_ti = pd.concat([df_ti, df_manual], ignore_index=True).drop_duplicates(subset=["Número", "MOVIMIENTO"], keep="last")
            except Exception as e:
                logger.warning(f"No se pudieron cargar novedades manuales: {e}")

        # Exportar a Excel con formato oficial y estilos de TI
        out_excel = self.outputs_dir / f"TI_Novedades_Estructura_Creacion_Emplea_{year}_{month:02d}.xlsx"
        
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Creación Emplea - Referidos"

        dark_fill = PatternFill(start_color="090D16", end_color="090D16", fill_type="solid")
        border_thin = Side(border_style="thin", color="CBD5E1")
        box_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

        # Write headers
        export_cols = [c for c in df_ti.columns if c != "picture_url"]
        ws.append(export_cols)
        for col_idx in range(1, len(export_cols) + 1):
            cell = ws.cell(1, col_idx)
            cell.fill = dark_fill
            cell.font = Font(name="Segoe UI", size=9, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center" if "Tipo" in str(cell.value) or "MOVIMIENTO" in str(cell.value) else "left")

        # Write rows
        for _, row in df_ti.iterrows():
            row_vals = [row[c] for c in export_cols]
            ws.append(row_vals)
            r_idx = ws.max_row
            for c_idx in range(1, len(export_cols) + 1):
                cell = ws.cell(r_idx, c_idx)
                cell.font = Font(name="Segoe UI", size=9)
                cell.border = box_border
                if export_cols[c_idx - 1] == "MOVIMIENTO":
                    mov = str(cell.value)
                    if mov == "INGRESO":
                        cell.fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
                        cell.font = Font(name="Segoe UI", size=9, bold=True, color="166534")
                    else:
                        cell.fill = PatternFill(start_color="FFE4E6", end_color="FFE4E6", fill_type="solid")
                        cell.font = Font(name="Segoe UI", size=9, bold=True, color="9F1239")
                    cell.alignment = Alignment(horizontal="center")

        # Adjust widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(out_excel)
        try:
            self.silver_dir.mkdir(parents=True, exist_ok=True)
            df_ti.to_parquet(self.silver_dir / f"ti_novedades_{year}_{month:02d}.parquet", index=False)
        except Exception:
            pass
        logger.info(f"Plantilla para TI generada con estilos oficiales: {len(df_ti)} novedades en {out_excel.name}")
        return df_ti


if __name__ == "__main__":
    p = SilverTINovedadesProcessor()
    df = p.generate_ti_template(2026, 7)
    print("Muestra Novedades TI:")
    print(df.head(6).to_string())
